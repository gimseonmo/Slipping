from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation

from app.models.presentation import PresentationState, SlideNotes

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


class PptxParseError(ValueError):
    pass


def parse_pptx(path: Path) -> PresentationState:
    if path.suffix.lower() != ".pptx":
        raise PptxParseError("PPTX 파일만 업로드할 수 있습니다.")
    if not zipfile.is_zipfile(path):
        raise PptxParseError("올바른 PPTX 파일이 아닙니다.")

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise PptxParseError(f"PPTX를 열 수 없습니다: {exc}") from exc

    notes_by_slide = _extract_notes_from_package(path)
    slides: list[SlideNotes] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text_runs: list[str] = []
        title = ""
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = "\n".join(
                paragraph.text.strip()
                for paragraph in shape.text_frame.paragraphs
                if paragraph.text.strip()
            ).strip()
            if not text:
                continue
            if not title:
                title = text.splitlines()[0]
            text_runs.append(text)

        slides.append(
            SlideNotes(
                slide_number=index,
                title=title or f"{index}번 슬라이드",
                slide_text="\n\n".join(text_runs),
                imported_notes=notes_by_slide.get(index, ""),
            )
        )

    if not slides:
        raise PptxParseError("슬라이드가 없는 PPTX입니다.")

    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    return PresentationState(
        presentation_id=digest,
        title=path.stem,
        source_filename=path.name,
        source_path=str(path),
        slide_count=len(slides),
        slides=slides,
        sync_message="발표 자료를 불러왔습니다",
    )


def _extract_notes_from_package(path: Path) -> dict[int, str]:
    notes_by_slide: dict[int, str] = {}
    with zipfile.ZipFile(path) as package:
        slide_names = sorted(
            (
                name
                for name in package.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ),
            key=_slide_number_from_name,
        )
        for slide_name in slide_names:
            slide_number = _slide_number_from_name(slide_name)
            rels_name = slide_name.replace("ppt/slides/", "ppt/slides/_rels/") + ".rels"
            if rels_name not in package.namelist():
                notes_by_slide[slide_number] = ""
                continue
            rels_root = ET.fromstring(package.read(rels_name))
            notes_target = None
            for rel in rels_root.findall("rel:Relationship", NS):
                rel_type = rel.attrib.get("Type", "")
                if rel_type.endswith("/notesSlide"):
                    notes_target = rel.attrib.get("Target")
                    break
            if not notes_target:
                notes_by_slide[slide_number] = ""
                continue
            notes_path = _resolve_package_path("ppt/slides", notes_target)
            if notes_path not in package.namelist():
                notes_by_slide[slide_number] = ""
                continue
            notes_by_slide[slide_number] = _extract_notes_text(package.read(notes_path))
    return notes_by_slide


def _extract_notes_text(xml_bytes: bytes) -> str:
    root = ET.fromstring(xml_bytes)
    candidate_paragraphs: list[str] = []
    fallback_paragraphs: list[str] = []

    for shape in root.findall(".//p:sp", NS):
        placeholder = shape.find(".//p:nvPr/p:ph", NS)
        placeholder_type = placeholder.attrib.get("type") if placeholder is not None else None
        paragraphs = _paragraphs_from_shape(shape)
        if not paragraphs:
            continue
        if placeholder_type in {"body", "notes"}:
            candidate_paragraphs.extend(paragraphs)
        elif placeholder_type not in {"hdr", "ftr", "dt", "sldNum"}:
            fallback_paragraphs.extend(paragraphs)

    paragraphs = candidate_paragraphs or fallback_paragraphs
    return "\n".join(paragraphs).strip()


def _paragraphs_from_shape(shape: ET.Element) -> list[str]:
    paragraphs: list[str] = []
    for paragraph in shape.findall(".//a:p", NS):
        pieces = [text.text or "" for text in paragraph.findall(".//a:t", NS)]
        line = "".join(pieces).strip()
        if line:
            paragraphs.append(line)
    return paragraphs


def _resolve_package_path(base_dir: str, target: str) -> str:
    base_parts = [part for part in base_dir.split("/") if part]
    for part in target.split("/"):
        if not part or part == ".":
            continue
        if part == "..":
            if base_parts:
                base_parts.pop()
        else:
            base_parts.append(part)
    return "/".join(base_parts)


def _slide_number_from_name(name: str) -> int:
    stem = Path(name).stem
    return int(stem.replace("slide", ""))
