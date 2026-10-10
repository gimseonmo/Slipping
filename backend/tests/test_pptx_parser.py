from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation

from app.services.pptx_parser import parse_pptx


def build_pptx_with_notes(path: Path) -> None:
    presentation = Presentation()
    for index in range(1, 4):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"{index}번 제목"
        slide.placeholders[1].text = f"{index}번 본문"
    presentation.save(path)

    with zipfile.ZipFile(path, "a") as package:
        for index in range(1, 4):
            notes_name = f"ppt/notesSlides/notesSlide{index}.xml"
            notes_text = "" if index == 2 else f"{index}번 발표자 노트\n둘째 줄"
            package.writestr(
                notes_name,
                f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:notes xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree>
    <p:sp>
      <p:nvSpPr><p:cNvPr id="2" name="Notes Placeholder"/><p:cNvSpPr/><p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr>
      <p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>{notes_text}</a:t></a:r></a:p></p:txBody>
    </p:sp>
  </p:spTree></p:cSld>
</p:notes>''',
            )
            rels_name = f"ppt/slides/_rels/slide{index}.xml.rels"
            original = package.read(rels_name).decode("utf-8")
            insert = f'<Relationship Id="rIdNotes{index}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide" Target="../notesSlides/notesSlide{index}.xml"/>'
            package.writestr(rels_name, original.replace("</Relationships>", f"{insert}</Relationships>"))


def test_parse_pptx_extracts_slide_text_and_notes(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    build_pptx_with_notes(path)

    state = parse_pptx(path)

    assert state.title == "sample"
    assert state.slide_count == 3
    assert state.slides[0].title == "1번 제목"
    assert "1번 본문" in state.slides[0].slide_text
    assert state.slides[0].imported_notes == "1번 발표자 노트\n둘째 줄"
    assert state.slides[1].imported_notes == ""
    assert state.slides[2].slide_number == 3
